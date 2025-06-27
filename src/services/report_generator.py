import os
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
import matplotlib.gridspec as gridspec
from datetime import datetime
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from src.utils.validators import ensure_required_columns


def add_overall_statistics(doc, df, output_dir, report_time, logger):
    """Add overall attrition statistics to the report"""
    try:
        # Debug: Log DataFrame columns and a sample row
        logger.info(f"DataFrame columns: {list(df.columns)}")
        logger.info(f"DataFrame head: {df.head(2).to_dict()}")
        doc.add_heading('1. Overall Attrition Statistics', level=1)
        # Calculate statistics
        if 'Action Type' not in df.columns or 'Employee Name' not in df.columns:
            logger.error("Required columns missing: 'Action Type' and/or 'Employee Name'")
            doc.add_paragraph("Error: Required columns missing: 'Action Type' and/or 'Employee Name'. Section skipped.")
            doc.add_page_break()
            return False
        total_employees = len(df)
        exits = df[df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)]
        total_exits = len(exits)
        attrition_rate = (total_exits / total_employees * 100).round(2) if total_employees > 0 else 0
        # Add statistics to document
        doc.add_paragraph(f"Total Employee Count: {total_employees}")
        doc.add_paragraph(f"Total Exits: {total_exits}")
        doc.add_paragraph(f"Overall Attrition Rate: {attrition_rate}%")
        # Create and add chart
        chart_path = create_overall_pie_chart(total_employees, total_exits, output_dir, report_time)
        doc.add_picture(chart_path, width=Inches(6))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        logger.info("Successfully added overall statistics section")
        return True
    except Exception as e:
        logger.error(f"Failed to add overall statistics section: {e}")
        logger.error(f"DataFrame columns at error: {list(df.columns)}")
        logger.error(f"DataFrame head at error: {df.head(2).to_dict()}")
        doc.add_paragraph("Error generating Overall Statistics section. Section skipped.")
        doc.add_page_break()
        return False

def add_gender_analysis(doc, df, output_dir, report_time, logger):
    """Add gender-wise attrition analysis to the report"""
    try:
        doc.add_heading('2. Gender-wise Attrition', level=1)
        attrition_by_gender = df[df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)].groupby('Gender').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'Attrition Count'})
        
        total_by_gender = df.groupby('Gender')['Employee Name'].count()
        attrition_by_gender['Total Employees'] = total_by_gender
        attrition_by_gender['Attrition Rate %'] = (attrition_by_gender['Attrition Count'] / attrition_by_gender['Total Employees'] * 100).round(2)
        
        # Add gender data table
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Gender'
        header_cells[1].text = 'Attrition Count'
        header_cells[2].text = 'Total Employees'
        header_cells[3].text = 'Attrition Rate %'

        style_table_headers(table)
        
        for index, row in attrition_by_gender.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Gender'])
            row_cells[1].text = str(row['Attrition Count'])
            row_cells[2].text = str(row['Total Employees'])
            row_cells[3].text = f"{row['Attrition Rate %']}%"
        
        # Create and add chart
        chart_path = create_gender_charts(attrition_by_gender, output_dir, report_time)
        doc.add_paragraph("\n")
        doc.add_picture(chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        
        logger.info("Successfully added gender analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add gender analysis section: {e}")
        doc.add_paragraph("Error generating Gender Analysis section. Section skipped.")
        doc.add_page_break()
        return False

def add_location_analysis(doc, df, output_dir, report_time, logger):
    """Add location-wise attrition analysis to the report"""
    try:
        doc.add_heading('3. Location-wise Attrition (Locations with ≥50 Employees)', level=1)
        
        total_by_location = df.groupby('Job Location')['Employee Name'].count()
        valid_locations = total_by_location[total_by_location >= 50].index
        
        attrition_by_location = df[
            (df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)) & 
            (df['Job Location'].isin(valid_locations))
        ].groupby('Job Location').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'Attrition Count'})
        
        total_by_location = total_by_location[valid_locations]
        attrition_by_location['Total Employees'] = total_by_location
        attrition_by_location['Attrition Rate %'] = (attrition_by_location['Attrition Count'] / attrition_by_location['Total Employees'] * 100).round(2)
        
        # Add location data table
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Location'
        header_cells[1].text = 'Attrition Count'
        header_cells[2].text = 'Total Employees'
        header_cells[3].text = 'Attrition Rate %'

        style_table_headers(table)
        
        for index, row in attrition_by_location.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Job Location'])
            row_cells[1].text = str(row['Attrition Count'])
            row_cells[2].text = str(row['Total Employees'])
            row_cells[3].text = f"{row['Attrition Rate %']}%"

        doc.add_paragraph("\n")
        
        # Create and add chart
        chart_path = create_location_charts(attrition_by_location, output_dir, report_time)
        doc.add_picture(chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_page_break()
        
        logger.info("Successfully added location analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add location analysis section: {e}")
        doc.add_paragraph("Error generating Location Analysis section. Section skipped.")
        doc.add_page_break()
        return False

def add_function_analysis(doc, df, output_dir, report_time, logger):
    """Add function-wise attrition analysis to the report"""
    try:
        doc.add_heading('4. Function-wise Attrition (Functions with ≥20 Employees)', level=1)
        
        total_by_function = df.groupby('Function')['Employee Name'].count()
        valid_functions = total_by_function[total_by_function >= 20].index
        
        attrition_by_function = df[
            (df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)) & 
            (df['Function'].isin(valid_functions))
        ].groupby('Function').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'Attrition Count'})
        
        total_by_function = total_by_function[valid_functions]
        attrition_by_function['Total Employees'] = total_by_function
        attrition_by_function['Attrition Rate %'] = (attrition_by_function['Attrition Count'] / attrition_by_function['Total Employees'] * 100).round(2)
        
        # Add function data table
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Function'
        header_cells[1].text = 'Attrition Count'
        header_cells[2].text = 'Total Employees'
        header_cells[3].text = 'Attrition Rate %'

        style_table_headers(table)
        
        for index, row in attrition_by_function.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Function'])
            row_cells[1].text = str(row['Attrition Count'])
            row_cells[2].text = str(row['Total Employees'])
            row_cells[3].text = f"{row['Attrition Rate %']}%"

        doc.add_paragraph("\n")
        
        # Create and add chart
        chart_path = create_function_charts(attrition_by_function, output_dir, report_time)
        doc.add_picture(chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_page_break()
        
        logger.info("Successfully added function analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add function analysis section: {e}")
        doc.add_paragraph("Error generating Function Analysis section. Section skipped.")
        doc.add_page_break()
        return False

def add_tenure_analysis(doc, df, output_dir, report_time, logger):
    """Add tenure analysis of exited employees to the report"""
    try:
        doc.add_heading('5. Tenure Analysis of Exited Employees', level=1)

        # Filter to only include employees who have exited
        exits_df = df[df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)].copy()

        # Calculate tenure in years (difference between joining date and action date)
        exits_df['Tenure in Days'] = (pd.to_datetime(exits_df['Action Date']) - pd.to_datetime(exits_df['Date of Joining'])).dt.days
        exits_df['Tenure in Years'] = exits_df['Tenure in Days'] / 365.25

        # Create tenure bins
        tenure_bins = [0, 1, 2, 3, 5, 10, float('inf')]
        tenure_labels = ['<1 year', '1-2 years', '2-3 years', '3-5 years', '5-10 years', '>10 years']
        exits_df['Tenure Band'] = pd.cut(exits_df['Tenure in Years'], bins=tenure_bins, labels=tenure_labels, right=False)

        # Group by tenure band - Fix for FutureWarning
        tenure_analysis = exits_df.groupby('Tenure Band', observed=True).size().reset_index(name='Count')
        tenure_analysis['Percentage'] = (tenure_analysis['Count'] / tenure_analysis['Count'].sum() * 100).round(2)

        # Add tenure data table
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Tenure Band'
        header_cells[1].text = 'Number of Exits'
        header_cells[2].text = 'Percentage'

        style_table_headers(table)

        for index, row in tenure_analysis.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Tenure Band'])
            row_cells[1].text = str(row['Count'])
            row_cells[2].text = f"{row['Percentage']}%"

        doc.add_paragraph("\n")
        
        # Create and add chart
        chart_path = create_tenure_charts(tenure_analysis, output_dir, report_time)
        doc.add_picture(chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break()
        
        logger.info("Successfully added tenure analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add tenure analysis section: {e}")
        doc.add_paragraph("Error generating Tenure Analysis section. Section skipped.")
        doc.add_page_break()
        return False

def add_grade_analysis(doc, df, output_dir, report_time, logger):
    """Add grade-wise attrition analysis to the report"""
    try:
        doc.add_heading('6. Grade-wise Attrition', level=1)

        # Get counts of employees by grade
        total_by_grade = df.groupby('Grade')['Employee Name'].count()
        valid_grades = total_by_grade[total_by_grade >= 10].index  # Only include grades with at least 10 employees

        # Filter to include only valid grades and exit actions
        attrition_by_grade = df[
            (df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)) & 
            (df['Grade'].isin(valid_grades))
        ].groupby('Grade').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'Attrition Count'})

        # Calculate attrition rates
        total_by_grade = total_by_grade[valid_grades]
        attrition_by_grade['Total Employees'] = total_by_grade
        attrition_by_grade['Attrition Rate %'] = (attrition_by_grade['Attrition Count'] / attrition_by_grade['Total Employees'] * 100).round(2)

        # Add grade data table
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Grade'
        header_cells[1].text = 'Attrition Count'
        header_cells[2].text = 'Total Employees'
        header_cells[3].text = 'Attrition Rate %'

        style_table_headers(table)
        
        for index, row in attrition_by_grade.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Grade'])
            row_cells[1].text = str(row['Attrition Count'])
            row_cells[2].text = str(row['Total Employees'])
            row_cells[3].text = f"{row['Attrition Rate %']}%"

        doc.add_paragraph("\n")
        
        # Create and add chart
        chart_path = create_grade_charts(attrition_by_grade, output_dir, report_time)
        doc.add_picture(chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_page_break()
        
        logger.info("Successfully added grade analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add grade analysis section: {e}")
        doc.add_paragraph("Error generating Grade Analysis section. Section skipped.")
        doc.add_page_break()
        return False

def add_trend_analysis(doc, df, output_dir, report_time, logger):
    """Add quarterly and monthly attrition trend analysis to the report"""
    try:
        doc.add_heading('7. Quarterly and Monthly Attrition Trends', level=1)

        # Filter to include only exit actions
        exits_df = df[df['Action Type'].str.contains('Exit|Resignation|Termination', na=False)].copy()

        # Ensure dates are in datetime format
        exits_df['Action Date'] = pd.to_datetime(exits_df['Action Date'])

        # Add month and quarter columns
        exits_df['Month'] = exits_df['Action Date'].dt.strftime('%Y-%m')
        exits_df['Quarter'] = exits_df['Action Date'].dt.to_period('Q').astype(str)
        
        # Calculate quarterly trends
        quarterly_trends = exits_df.groupby('Quarter').size().reset_index(name='Exit Count')
        quarterly_trends = quarterly_trends.sort_values('Quarter')  # Sort by quarter chronologically
        quarterly_trends['Quarter'] = pd.PeriodIndex(quarterly_trends['Quarter'], freq='Q')

        # Calculate monthly trends
        monthly_trends = exits_df.groupby('Month').size().reset_index(name='Exit Count')
        monthly_trends = monthly_trends.sort_values('Month')  # Sort by month chronologically
        monthly_trends['Month'] = pd.to_datetime(monthly_trends['Month'])

        # Add quarterly trend table first
        doc.add_heading('Quarterly Attrition Trend', level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Quarter'
        header_cells[1].text = 'Exit Count'

        style_table_headers(table)

        for index, row in quarterly_trends.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Quarter'])
            row_cells[1].text = str(row['Exit Count'])

        doc.add_paragraph("\n")
        
        # Create and add quarterly trend chart
        quarterly_chart_path = create_quarterly_trend_chart(quarterly_trends, output_dir, report_time)
        doc.add_picture(quarterly_chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Now add monthly trend table
        doc.add_heading('Monthly Attrition Trend', level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Month'
        header_cells[1].text = 'Exit Count'

        style_table_headers(table)
        
        for index, row in monthly_trends.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = row['Month'].strftime('%Y-%m')
            row_cells[1].text = str(row['Exit Count'])

        doc.add_paragraph("\n")

        # Create and add monthly trend chart
        monthly_chart_path = create_monthly_trend_chart(monthly_trends, output_dir, report_time)
        doc.add_picture(monthly_chart_path, width=Inches(7))
        
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        logger.info("Successfully added trend analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add trend analysis section: {e}")
        doc.add_paragraph("Error generating Trend Analysis section. Section skipped.")
        return False

def generate_attrition_report(df, output_dir, logger):
    """
    Create a comprehensive automated attrition analysis report
    """
    # Create document
    doc, report_time = create_report_document(output_dir)
    
    if doc is None:
        logger.error("Failed to create report document. Aborting report generation.")
        return False, None, None
    
    # Define all report sections to run
    report_sections = [
        {"func": add_overall_statistics, "name": "Overall Statistics"},
        {"func": add_gender_analysis, "name": "Gender Analysis"},
        {"func": add_location_analysis, "name": "Location Analysis"},
        {"func": add_function_analysis, "name": "Function Analysis"},
        {"func": add_tenure_analysis, "name": "Tenure Analysis"},
        {"func": add_grade_analysis, "name": "Grade Analysis"},
        {"func": add_trend_analysis, "name": "Trend Analysis"}
    ]
    
    # Execute each section
    successful_sections = 0
    
    for section in report_sections:
        section_func = section["func"]
        section_name = section["name"]
        
        logger.info(f"Generating {section_name} section...")
        
        try:
            result = section_func(doc, df, output_dir, report_time, logger)
            if result:
                successful_sections += 1
            else:
                logger.warning(f"Section function returned False: {section_name}")
        except Exception as e:
            logger.error(f"Error in {section_name} section: {e}")
            doc.add_paragraph(f"Error generating {section_name} section. Section skipped.")
            doc.add_page_break()
    
    # Save the document
    try:
        report_path = save_document(doc, output_dir, report_time)
        logger.info(f"Report saved to {report_path}")
        logger.info(f"Report generated with {successful_sections} of {len(report_sections)} sections completed successfully")
        return True, report_path, report_time
    except Exception as e:
        logger.error(f"Failed to save report: {e}")
        return False, None, None
    
def create_overall_pie_chart(total_employees, total_exits, output_dir, report_time):
    """Create overall attrition pie chart"""
    plt.figure(figsize=(10, 6))
    labels = ['Active Employees', 'Exited Employees']
    values = [total_employees - total_exits, total_exits]
    plt.pie(values, labels=labels, autopct='%1.1f%%', startangle=90, explode=[0, 0.1])
    plt.title('Overall Attrition', fontsize=16, fontweight='bold')
    plt.tight_layout()
    
    chart_path = f"{output_dir}/Overall_Statistics_{report_time}.png"
    plt.savefig(chart_path)
    plt.close()
    
    return chart_path

def create_gender_charts(attrition_by_gender, output_dir, report_time):
    """Create gender analysis charts"""
    fig = plt.figure(figsize=(16, 7))
    gs = gridspec.GridSpec(1, 2, width_ratios=[2, 3])
    ax1 = plt.subplot(gs[0])
    ax2 = plt.subplot(gs[1])

    sns.barplot(x=attrition_by_gender.index, y=attrition_by_gender['Attrition Count'], 
                ax=ax1, hue=attrition_by_gender.index, palette="coolwarm", width=0.4)
    ax1.set_title('Attrition Count by Gender\n', fontsize=15, fontweight='bold', color='black')
    ax1.set_ylabel('Number of Employees', fontsize=14)
    ax1.set_xlabel('Gender', fontsize=14)
    
    ax2.pie(attrition_by_gender['Attrition Rate %'], labels=attrition_by_gender.index, 
            autopct='%1.1f%%', textprops={'fontsize': 14})
    ax2.set_title('Attrition Rate Distribution by Gender\n', fontsize=15, fontweight='bold', color='black')
    plt.tight_layout()
    
    chart_path = f"{output_dir}/Gender_Analysis_{report_time}.png"
    plt.savefig(chart_path)
    plt.close()
    
    return chart_path

def create_location_charts(attrition_by_location, output_dir, report_time):
    """Create location analysis charts"""
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7))
    sns.barplot(data=attrition_by_location.reset_index(), x='Job Location', y='Attrition Count', ax=ax1)
    ax1.set_title('Attrition Count by Location\n', fontsize=16, fontweight='bold')
    ax1.set_xlabel('Location', fontsize=14)
    ax1.set_ylabel('Number of Employees', fontsize=14)
    ax1.tick_params(axis='x', rotation=45)
    
    ax2.pie(attrition_by_location['Attrition Rate %'], labels=attrition_by_location.index, 
            autopct='%1.1f%%', textprops={'fontsize': 14})
    ax2.set_title('Attrition Rate Distribution by Location\n', fontsize=16, fontweight='bold')
    plt.tight_layout()
    
    chart_path = f"{output_dir}/Location_Analysis_{report_time}.png"
    plt.savefig(chart_path, bbox_inches='tight')
    plt.close()
    
    return chart_path

def create_function_charts(attrition_by_function, output_dir, report_time):
    """Create function analysis charts"""
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(17, 7))
    sns.barplot(data=attrition_by_function.reset_index(), x='Function', y='Attrition Count', 
                hue='Function', ax=ax1, legend=False)
    ax1.set_title('Attrition Count by Function\n', fontsize=16, fontweight='bold')
    ax1.set_ylabel('Number of Employees', fontsize=14)
    ax1.set_xlabel('Function', fontsize=14)
    ax1.tick_params(axis='x', rotation=45)
    
    ax2.pie(attrition_by_function['Attrition Rate %'], labels=attrition_by_function.index, 
            autopct='%1.1f%%', textprops={'fontsize': 14})
    ax2.set_title('Attrition Rate Distribution by Function\n', fontsize=16, fontweight='bold')
    plt.tight_layout()
    
    chart_path = f"{output_dir}/Function_Analysis_{report_time}.png"
    plt.savefig(chart_path, bbox_inches='tight')
    plt.close()
    
    return chart_path

def create_tenure_charts(tenure_analysis, output_dir, report_time):
    """Create tenure analysis charts"""
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7))

    # Bar chart - Fix for seaborn FutureWarning
    sns.barplot(data=tenure_analysis, x='Tenure Band', y='Count', hue='Tenure Band', ax=ax1, legend=False)
    ax1.set_title('Attrition Count by Tenure\n', fontsize=16, fontweight='bold')
    ax1.set_ylabel('Number of Employees', fontsize=14)
    ax1.set_xlabel('Tenure Band', fontsize=14)
    ax1.tick_params(axis='x', rotation=45)

    # Pie chart
    ax2.pie(tenure_analysis['Percentage'], labels=tenure_analysis['Tenure Band'], 
            autopct='%1.1f%%', textprops={'fontsize': 14})
    ax2.set_title('Attrition Distribution by Tenure\n', fontsize=16, fontweight='bold')
    plt.tight_layout()

    chart_path = f"{output_dir}/Tenure_Analysis_{report_time}.png"
    plt.savefig(chart_path, bbox_inches='tight')
    plt.close()
    
    return chart_path

def create_grade_charts(attrition_by_grade, output_dir, report_time):
    """Create grade analysis charts"""
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7))

    # Bar chart
    sns.barplot(data=attrition_by_grade.reset_index(), x='Grade', y='Attrition Count', 
                hue='Grade', ax=ax1, legend=False)
    ax1.set_title('Attrition Count by Grade\n', fontsize=16, fontweight='bold')
    ax1.set_ylabel('Number of Employees', fontsize=14)
    ax1.set_xlabel('Grade', fontsize=14)
    ax1.tick_params(axis='x', rotation=45)

    # Pie chart
    ax2.pie(attrition_by_grade['Attrition Rate %'], labels=attrition_by_grade.index, 
            autopct='%1.1f%%', textprops={'fontsize': 14})
    ax2.set_title('Attrition Rate Distribution by Grade\n', fontsize=16, fontweight='bold')
    plt.tight_layout()

    chart_path = f"{output_dir}/Grade_Analysis_{report_time}.png"
    plt.savefig(chart_path, bbox_inches='tight')
    plt.close()
    
    return chart_path

def create_quarterly_trend_chart(quarterly_trends, output_dir, report_time):
    """Create quarterly trend chart"""
    plt.figure(figsize=(10, 6))
    ax = sns.barplot(data=quarterly_trends, x=quarterly_trends['Quarter'].astype(str), 
                     y='Exit Count', hue=quarterly_trends['Quarter'].astype(str), legend=False)
    
    plt.title('Quarterly Attrition Trend\n', fontsize=16, fontweight='bold')
    plt.ylabel('Number of Exits', fontsize=14)
    plt.xlabel('Quarter', fontsize=14)
    plt.grid(True, linestyle='--', alpha=0.7, axis='y')
    plt.tight_layout()

    # Add count labels on top of each bar
    for i, v in enumerate(quarterly_trends['Exit Count']):
        ax.text(i, v + 0.5, str(v), ha='center', fontweight='bold')

    chart_path = f"{output_dir}/Quarterly_Trend_{report_time}.png"
    plt.savefig(chart_path, bbox_inches='tight')
    plt.close()
    
    return chart_path

def create_monthly_trend_chart(monthly_trends, output_dir, report_time):
    """Create monthly trend chart"""
    plt.figure(figsize=(12, 6))
    sns.lineplot(data=monthly_trends, x='Month', y='Exit Count', marker='o', linewidth=2)
    plt.title('Monthly Attrition Trend\n', fontsize=16, fontweight='bold')
    plt.ylabel('Number of Exits', fontsize=14)
    plt.xlabel('Month', fontsize=14)
    plt.xticks(rotation=45)
    plt.grid(True, linestyle='--', alpha=0.7)
    plt.tight_layout()

    chart_path = f"{output_dir}/Monthly_Trend_{report_time}.png"
    plt.savefig(chart_path, bbox_inches='tight')
    plt.close()
    
    return chart_path    
def load_data(file_path: str, logger):
    """
    Load and prepare the data
    
    Args:
        file_path: Path to the Excel file
        logger: Logger instance for logging operations
        
    Returns:
        DataFrame if successful, None if error occurs
    """
    try:
        df = pd.read_excel(file_path)
        
        # Ensure all required columns exist
        ensure_required_columns(df)
        
        # Fill missing values in Action Type column
        df['Action Type'] = df['Action Type'].fillna('')
        
        logger.info(f"Successfully loaded data from {file_path}")
        return df
    except Exception as e:
        logger.error(f"Error loading data: {e}")
        return None


def create_report_document(output_dir):
    """
    Create and initialize the Word document for the report
    """
    try:
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        # Create Word document
        doc = Document()
        
        # Add title
        doc.add_heading('AUTOMATED ATTRITION ANALYSIS REPORT BY VIMAL SINGH', 0)
        doc.add_paragraph(f"Report Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        
        return doc, datetime.now().strftime("%Y%m%d_%H%M%S")
    except Exception as e:
        return None, None

def style_table_headers(table, header_color="ADD8E6"):
    """
    Apply styling to the header row of a table
    
    Parameters:
    - table: the table object to style
    - header_color: hex color code for header background (default: light blue)
    """
    try:
        header_cells = table.rows[0].cells
        
        for cell in header_cells:
            # Make text bold
            if cell.paragraphs and cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].bold = True
            # Center text
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Apply background color
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{header_color}"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
        return True
    except Exception as e:
        return False

def save_document(doc, output_dir, report_time):
    """Save the Word document"""
    today = datetime.now().strftime("%Y-%m-%d")
    time = datetime.now().strftime("%H%M")  # Removed colon for file naming compatibility
    report_path = f"{output_dir}/Attrition_Report_{today}_{time}.docx"
    doc.save(report_path)
    return report_path