from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from datetime import datetime
import os

def create_report_document(output_dir):
    """
    Create and initialize the Word document for the report
    """
    try:
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        # Create Word document
        doc = Document()
        
        # Add title
        doc.add_heading('AUTOMATED ATTRITION ANALYSIS REPORT BY VIMAL SINGH', 0)
        doc.add_paragraph(f"Report Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        
        return doc, datetime.now().strftime("%Y%m%d_%H%M%S")
    except Exception as e:
        return None, None

def style_table_headers(table, header_color="ADD8E6"):
    """
    Apply styling to the header row of a table
    
    Parameters:
    - table: the table object to style
    - header_color: hex color code for header background (default: light blue)
    """
    try:
        header_cells = table.rows[0].cells
        
        for cell in header_cells:
            # Make text bold
            if cell.paragraphs and cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].bold = True
            # Center text
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Apply background color
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{header_color}"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
        return True
    except Exception as e:
        return False

def save_document(doc, output_dir, report_time):
    """Save the Word document"""
    today = datetime.now().strftime("%Y-%m-%d")
    time = datetime.now().strftime("%H%M")  # Removed colon for file naming compatibility
    report_path = f"{output_dir}/Attrition_Report_{today}_{time}.docx"
    doc.save(report_path)
    return report_path