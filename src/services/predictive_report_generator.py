import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from datetime import datetime
import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import logging
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestClassifier
from sklearn.preprocessing import LabelEncoder
import io
import tempfile
from pathlib import Path

from src.core.logging import setup_logging

# Initialize logging
logger = setup_logging()


def load_data(file_content: bytes, filename: str) -> pd.DataFrame:
    """Load and validate Excel data."""
    try:
        df = pd.read_excel(io.BytesIO(file_content))
        required_columns = ['Employee Name', 'Gender', 'Function', 'Job Location', 'Grade', 'Action Date', 'Date of Joining', 'Action Type']
        missing_columns = [col for col in required_columns if col not in df.columns]
        if missing_columns:
            raise ValueError(f"Missing required columns: {', '.join(missing_columns)}")
        
        # Validate date columns
        for col in ['Action Date', 'Date of Joining']:
            try:
                df[col] = pd.to_datetime(df[col], errors='coerce')
                if df[col].isna().any():
                    raise ValueError(f"Invalid or missing dates in {col}")
            except Exception as e:
                raise ValueError(f"Error parsing {col}: {str(e)}")
        
        df['Action Type'] = df['Action Type'].fillna('')
        logger.info(f"Successfully loaded data from {filename}. Rows: {len(df)}")
        return df
    except Exception as e:
        logger.error(f"Error loading data from {filename}: {e}")
        raise

def style_table_headers(table, header_color="ADD8E6") -> bool:
    """Apply styling to table headers."""
    try:
        header_cells = table.rows[0].cells
        for cell in header_cells:
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{header_color}"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
        return True
    except Exception as e:
        logger.warning(f"Failed to style table headers: {e}")
        return False

def create_report_document() -> tuple[Document, str]:
    """Initialize Word document for the report."""
    try:
        doc = Document()
        doc.add_heading('PREDICTIVE ANALYTICS REPORT: ATTRITION RISK', 0)
        doc.add_paragraph(f"Report Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        return doc, datetime.now().strftime("%Y%m%d_%H%M%S")
    except Exception as e:
        logger.error(f"Failed to create report document: {e}")
        raise

def add_predictive_analytics(doc: Document, df: pd.DataFrame, output_dir: str, report_time: str, model: RandomForestClassifier) -> tuple[pd.DataFrame, bool]:
    """Add predictive analytics section with top 50 high-risk employees."""
    try:
        doc.add_heading('1. Predictive Analytics for Attrition Risk', level=1)
        doc.add_paragraph("Lists top 50 active employees with attrition probability > 70%, sorted by probability, based on a Random Forest Classifier.")

        # Prepare data
        df['Is_Separation'] = df['Action Type'].str.contains('Separation|Termination', case=False, na=False).astype(int)
        features = ['Gender', 'Tenure in Years', 'Function', 'Grade']
        df['Tenure in Years'] = (pd.to_datetime(df['Action Date']) - pd.to_datetime(df['Date of Joining'])).dt.days / 365.25
        if df['Tenure in Years'].isna().any():
            raise ValueError("Tenure calculation failed due to invalid dates")
        
        X = df[features].copy()
        y = df['Is_Separation']

        # Encode categorical variables
        encoders = {}
        for col in X.select_dtypes(include=['object']).columns:
            X[col] = X[col].astype(str)
            encoders[col] = LabelEncoder()
            X[col] = encoders[col].fit_transform(X[col])
        X = X.fillna(0)

        # Train model
        X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)
        model.fit(X_train, y_train)

        # Predict for active employees
        # Log Action Type distribution for debugging
        action_types = df['Action Type'].value_counts(dropna=False).to_dict()
        logger.info(f"Action Type distribution: {action_types}")
        logger.info(f"Total rows before filtering: {len(df)}")
        
        active_df = df[~df['Action Type'].str.contains('Separation|Termination', case=False, na=False)].copy()
        logger.info(f"Active employees after filtering: {len(active_df)}")

        if active_df.empty:
            logger.warning("No active employees found. Falling back to include all employees with empty or unknown Action Type.")
            # Fallback: Treat employees with empty or unknown Action Type as active
            active_df = df[df['Action Type'].str.strip() == ''].copy()
            logger.info(f"Active employees after fallback: {len(active_df)}")
            if active_df.empty:
                logger.warning("No employees with empty Action Type found. No predictions possible.")
                doc.add_paragraph("No active employees available for prediction.")
                return None, True

        X_active = active_df[features].copy()
        for col in X_active.select_dtypes(include=['object']).columns:
            X_active[col] = X_active[col].astype(str)
            X_active[col] = encoders[col].transform(X_active[col])
        X_active = X_active.fillna(0)
        active_df['Attrition_Probability'] = model.predict_proba(X_active)[:, 1]
        high_risk_employees = active_df[active_df['Attrition_Probability'] > 0.7][['Employee Name', 'Gender', 'Function', 'Job Location', 'Grade', 'Action Date', 'Date of Joining', 'Attrition_Probability']]

        # Add total high-risk count
        total_high_risk = len(high_risk_employees)
        doc.add_paragraph(f"Total Number of High-Risk Employees: {total_high_risk}")
        logger.info(f"High-risk employees: {total_high_risk}")

        # Add table for top 50
        high_risk_display = high_risk_employees.sort_values(by='Attrition_Probability', ascending=False).head(50)
        doc.add_paragraph("Top 50 Employees with High Attrition Risk (Probability > 0.7):")
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Employee Name'
        header_cells[1].text = 'Function'
        header_cells[2].text = 'Attrition Probability'
        style_table_headers(table)

        if high_risk_display.empty:
            logger.warning("No high-risk employees found.")
            doc.add_paragraph("No active employees with attrition probability > 0.7.")
        else:
            for _, row in high_risk_display.iterrows():
                row_cells = table.add_row().cells
                row_cells[0].text = str(row['Employee Name'])
                row_cells[1].text = str(row['Function'])
                row_cells[2].text = f"{row['Attrition_Probability']:.2f}"

        doc.add_page_break()
        logger.info("Added predictive analytics section")
        return high_risk_employees, True
    except Exception as e:
        logger.error(f"Failed to add predictive analytics: {e}")
        doc.add_paragraph(f"Error generating Predictive Analytics section: {str(e)}")
        doc.add_page_break()
        return None, False

def add_gender_analysis(doc: Document, high_risk_df: pd.DataFrame, output_dir: str, report_time: str) -> bool:
    """Add gender-wise analysis of high-risk employees."""
    try:
        doc.add_heading('2. Gender-wise Attrition Risk Analysis', level=1)
        if high_risk_df is None or high_risk_df.empty:
            doc.add_paragraph("No high-risk employees available.")
            return True

        attrition_by_gender = high_risk_df.groupby('Gender').agg({'Employee Name': 'count'}).rename(columns={'Employee Name': 'High-Risk Count'})
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Gender'
        header_cells[1].text = 'High-Risk Count'
        style_table_headers(table)

        for index, row in attrition_by_gender.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Gender'])
            row_cells[1].text = str(row['High-Risk Count'])

        fig = plt.figure(figsize=(8, 6))
        sns.barplot(x=attrition_by_gender.index, y=attrition_by_gender['High-Risk Count'], hue=attrition_by_gender.index, palette="coolwarm")
        plt.title('High-Risk Count by Gender\n', fontsize=15, fontweight='bold')
        plt.ylabel('Number of Employees', fontsize=14)
        plt.xlabel('Gender', fontsize=14)
        plt.tight_layout()
        gender_plot_path = f"{output_dir}/image1_{report_time}.png"
        plt.savefig(gender_plot_path)
        plt.close()

        doc.add_paragraph("\n")
        doc.add_picture(gender_plot_path, width=Inches(5))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        logger.info("Added gender analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add gender analysis: {e}")
        doc.add_paragraph("Error generating Gender Analysis section.")
        doc.add_page_break()
        return False

def add_location_analysis(doc: Document, high_risk_df: pd.DataFrame, output_dir: str, report_time: str) -> bool:
    """Add location-wise analysis of high-risk employees."""
    try:
        doc.add_heading('3. Location-wise Attrition Risk Analysis', level=1)
        if high_risk_df is None or high_risk_df.empty:
            doc.add_paragraph("No high-risk employees available.")
            return True

        total_by_location = high_risk_df.groupby('Job Location')['Employee Name'].count()
        valid_locations = total_by_location[total_by_location > 2].index
        if valid_locations.empty:
            doc.add_paragraph("No locations with more than 2 high-risk employees.")
            return True

        attrition_by_location = high_risk_df[high_risk_df['Job Location'].isin(valid_locations)].groupby('Job Location').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'High-Risk Count'})

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Location'
        header_cells[1].text = 'High-Risk Count'
        style_table_headers(table)

        for index, row in attrition_by_location.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Job Location'])
            row_cells[1].text = str(row['High-Risk Count'])

        fig, ax = plt.subplots(figsize=(10, 6))
        sns.barplot(data=attrition_by_location.reset_index(), x='Job Location', y='High-Risk Count', hue='Job Location', ax=ax)
        ax.set_title('High-Risk Count by Location\n', fontsize=16, fontweight='bold')
        ax.set_xlabel('Location', fontsize=14)
        ax.set_ylabel('Number of Employees', fontsize=14)
        ax.tick_params(axis='x', rotation=45)
        plt.tight_layout()
        location_plot_path = f"{output_dir}/image2_{report_time}.png"
        plt.savefig(location_plot_path, bbox_inches='tight')
        plt.close()

        doc.add_paragraph("\n")
        doc.add_picture(location_plot_path, width=Inches(6))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        logger.info("Added location analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add location analysis: {e}")
        doc.add_paragraph("Error generating Location Analysis section.")
        doc.add_page_break()
        return False

def add_function_analysis(doc: Document, high_risk_df: pd.DataFrame, output_dir: str, report_time: str) -> bool:
    """Add function-wise analysis of high-risk employees."""
    try:
        doc.add_heading('4. Function-wise Attrition Risk Analysis', level=1)
        if high_risk_df is None or high_risk_df.empty:
            doc.add_paragraph("No high-risk employees available.")
            return True

        total_by_function = high_risk_df.groupby('Function')['Employee Name'].count()
        valid_functions = total_by_function[total_by_function > 1].index
        if valid_functions.empty:
            doc.add_paragraph("No functions with more than 1 high-risk employee.")
            return True

        attrition_by_function = high_risk_df[high_risk_df['Function'].isin(valid_functions)].groupby('Function').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'High-Risk Count'})

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Function'
        header_cells[1].text = 'High-Risk Count'
        style_table_headers(table)

        for index, row in attrition_by_function.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Function'])
            row_cells[1].text = str(row['High-Risk Count'])

        fig, ax = plt.subplots(figsize=(10, 6))
        sns.barplot(data=attrition_by_function.reset_index(), x='Function', y='High-Risk Count', hue='Function', ax=ax)
        ax.set_title('High-Risk Count by Function\n', fontsize=16, fontweight='bold')
        ax.set_ylabel('Number of Employees', fontsize=14)
        ax.set_xlabel('Function', fontsize=14)
        ax.tick_params(axis='x', rotation=45)
        plt.tight_layout()
        function_plot_path = f"{output_dir}/image3_{report_time}.png"
        plt.savefig(function_plot_path, bbox_inches='tight')
        plt.close()

        doc.add_paragraph("\n")
        doc.add_picture(function_plot_path, width=Inches(6))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        logger.info("Added function analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add function analysis: {e}")
        doc.add_paragraph("Error generating Function Analysis section.")
        doc.add_page_break()
        return False

def add_tenure_analysis(doc: Document, high_risk_df: pd.DataFrame, output_dir: str, report_time: str) -> bool:
    """Add tenure analysis of high-risk employees."""
    try:
        doc.add_heading('5. Tenure Analysis of High-Risk Employees', level=1)
        if high_risk_df is None or high_risk_df.empty:
            doc.add_paragraph("No high-risk employees available.")
            return True

        high_risk_df = high_risk_df.copy()
        high_risk_df['Tenure in Days'] = (pd.to_datetime(high_risk_df['Action Date']) - pd.to_datetime(high_risk_df['Date of Joining'])).dt.days
        high_risk_df['Tenure in Years'] = high_risk_df['Tenure in Days'] / 365.25
        tenure_bins = [0, 1, 2, 3, 5, 10, float('inf')]
        tenure_labels = ['<1 year', '1-2 years', '2-3 years', '3-5 years', '5-10 years', '>10 years']
        high_risk_df['Tenure Band'] = pd.cut(high_risk_df['Tenure in Years'], bins=tenure_bins, labels=tenure_labels, right=False)

        tenure_analysis = high_risk_df.groupby('Tenure Band', observed=True).size().reset_index(name='High-Risk Count')
        tenure_analysis = tenure_analysis[tenure_analysis['High-Risk Count'] > 2]
        if tenure_analysis.empty:
            doc.add_paragraph("No tenure bands with more than 2 high-risk employees.")
            return True

        tenure_analysis['Percentage'] = (tenure_analysis['High-Risk Count'] / tenure_analysis['High-Risk Count'].sum() * 100).round(2)
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Tenure Band'
        header_cells[1].text = 'High-Risk Count'
        header_cells[2].text = 'Percentage'
        style_table_headers(table)

        for index, row in tenure_analysis.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Tenure Band'])
            row_cells[1].text = str(row['High-Risk Count'])
            row_cells[2].text = f"{row['Percentage']}%"

        fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 6))
        sns.barplot(data=tenure_analysis, x='Tenure Band', y='High-Risk Count', hue='Tenure Band', ax=ax1, legend=False)
        ax1.set_title('High-Risk Count by Tenure\n', fontsize=16, fontweight='bold')
        ax1.set_ylabel('Number of Employees', fontsize=14)
        ax1.set_xlabel('Tenure Band', fontsize=14)
        ax1.tick_params(axis='x', rotation=45)
        ax2.pie(tenure_analysis['Percentage'], labels=tenure_analysis['Tenure Band'], autopct='%1.1f%%', textprops={'fontsize': 14})
        ax2.set_title('High-Risk Distribution by Tenure\n', fontsize=16, fontweight='bold')
        plt.tight_layout()
        tenure_plot_path = f"{output_dir}/image4_{report_time}.png"
        plt.savefig(tenure_plot_path, bbox_inches='tight')
        plt.close()

        doc.add_paragraph("\n")
        doc.add_picture(tenure_plot_path, width=Inches(6))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        logger.info("Added tenure analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add tenure analysis: {e}")
        doc.add_paragraph("Error generating Tenure Analysis section.")
        doc.add_page_break()
        return False

def add_grade_analysis(doc: Document, high_risk_df: pd.DataFrame, output_dir: str, report_time: str) -> bool:
    """Add grade-wise analysis of high-risk employees."""
    try:
        doc.add_heading('6. Grade-wise Attrition Risk Analysis', level=1)
        if high_risk_df is None or high_risk_df.empty:
            doc.add_paragraph("No high-risk employees available.")
            return True

        total_by_grade = high_risk_df.groupby('Grade')['Employee Name'].count()
        valid_grades = total_by_grade[total_by_grade > 2].index
        if valid_grades.empty:
            doc.add_paragraph("No grades with more than 2 high-risk employees.")
            return True

        attrition_by_grade = high_risk_df[high_risk_df['Grade'].isin(valid_grades)].groupby('Grade').agg({
            'Employee Name': 'count'
        }).rename(columns={'Employee Name': 'High-Risk Count'})

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Grade'
        header_cells[1].text = 'High-Risk Count'
        style_table_headers(table)

        for index, row in attrition_by_grade.reset_index().iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Grade'])
            row_cells[1].text = str(row['High-Risk Count'])

        fig, ax = plt.subplots(figsize=(10, 6))
        sns.barplot(data=attrition_by_grade.reset_index(), x='Grade', y='High-Risk Count', hue='Grade', ax=ax)
        ax.set_title('High-Risk Count by Grade\n', fontsize=16, fontweight='bold')
        ax.set_ylabel('Number of Employees', fontsize=14)
        ax.set_xlabel('Grade', fontsize=14)
        ax.tick_params(axis='x', rotation=45)
        plt.tight_layout()
        grade_plot_path = f"{output_dir}/image5_{report_time}.png"
        plt.savefig(grade_plot_path, bbox_inches='tight')
        plt.close()

        doc.add_paragraph("\n")
        doc.add_picture(grade_plot_path, width=Inches(6))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        logger.info("Added grade analysis section")
        return True
    except Exception as e:
        logger.error(f"Failed to add grade analysis: {e}")
        doc.add_paragraph("Error generating Grade Analysis section.")
        doc.add_page_break()
        return False

def add_complete_high_risk_list(doc: Document, high_risk_df: pd.DataFrame, output_dir: str, report_time: str) -> bool:
    """Add complete list of high-risk employees."""
    try:
        doc.add_heading('7. Complete List of High-Risk Employees', level=1)
        doc.add_paragraph("Complete list of active employees with attrition probability > 70%, sorted by probability.")

        if high_risk_df is None or high_risk_df.empty:
            doc.add_paragraph("No high-risk employees available.")
            return True

        high_risk_df = high_risk_df.sort_values(by='Attrition_Probability', ascending=False)
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Employee Name'
        header_cells[1].text = 'Function'
        header_cells[2].text = 'Attrition Probability'
        style_table_headers(table)

        for _, row in high_risk_df.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['Employee Name'])
            row_cells[1].text = str(row['Function'])
            row_cells[2].text = f"{row['Attrition_Probability']:.2f}"

        doc.add_page_break()
        logger.info("Added complete high-risk list section")
        return True
    except Exception as e:
        logger.error(f"Failed to add complete high-risk list: {e}")
        doc.add_paragraph("Error generating Complete High-Risk List section.")
        doc.add_page_break()
        return False

def add_model_details(doc: Document, model: RandomForestClassifier, features: list, output_dir: str, report_time: str) -> bool:
    """Add feature importance table for the predictive model."""
    try:
        doc.add_heading('8. Model Details', level=1)
        doc.add_paragraph("The table below shows the feature importance of the Random Forest Classifier used to predict attrition risk:")

        feature_importance = [
            {'Feature': feature, 'Importance': importance}
            for feature, importance in zip(features, model.feature_importances_)
        ]
        feature_importance.sort(key=lambda x: x['Importance'], reverse=True)

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Feature'
        header_cells[1].text = 'Importance'
        style_table_headers(table)

        for item in feature_importance:
            row_cells = table.add_row().cells
            row_cells[0].text = item['Feature']
            row_cells[1].text = f"{item['Importance']:.4f}"

        doc.add_page_break()
        logger.info("Added model details section")
        return True
    except Exception as e:
        logger.error(f"Failed to add model details: {e}")
        doc.add_paragraph("Error generating Model Details section.")
        doc.add_page_break()
        return False

def create_predictive_report(df: pd.DataFrame, output_dir: str) -> str:
    """Generate predictive attrition report."""
    try:
        doc, report_time = create_report_document()
        model = RandomForestClassifier(random_state=42)
        features = ['Gender', 'Tenure in Years', 'Function', 'Grade']
        high_risk_df, success = add_predictive_analytics(doc, df, output_dir, report_time, model)
        if success:
            report_sections = [
                (add_gender_analysis, "Gender Analysis"),
                (add_location_analysis, "Location Analysis"),
                (add_function_analysis, "Function Analysis"),
                (add_tenure_analysis, "Tenure Analysis"),
                (add_grade_analysis, "Grade Analysis"),
                (add_complete_high_risk_list, "Complete High-Risk List"),
                (lambda doc, df, out, time: add_model_details(doc, model, features, out, time), "Model Details")
            ]
            for section_func, section_name in report_sections:
                logger.info(f"Generating {section_name} section...")
                section_func(doc, high_risk_df, output_dir, report_time)

        report_path = f"{output_dir}/Attrition_Predictive_Report_{report_time}.docx"
        doc.save(report_path)
        logger.info(f"Report saved to {report_path}")
        return report_path
    except Exception as e:
        logger.error(f"Failed to save report: {e}")
        raise